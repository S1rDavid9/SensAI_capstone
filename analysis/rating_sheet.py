"""Builds the decision-plausibility rating spreadsheet from the blind sheet
written by analysis/stage2_on_pilot.py, and scores it once rated.

    uv run python -m analysis.rating_sheet build    # -> pilot/data/plausibility_rating_sheet.xlsx
    uv run python -m analysis.rating_sheet score    # reads the ratings back -> analysis/results/plausibility.json

The spreadsheet holds learner-level rows, so it stays in pilot/data/ (never
committed). Raters do not see which action came from the DQN alone and
which after a guardrail; that key is kept in a separate file.
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from pilot.core.tasks import PILOT_FORM_IDS, QuestionBank

ROOT = Path(__file__).parent.parent
DATA = ROOT / "pilot" / "data"
SHEET_CSV = DATA / "plausibility_rating_sheet.csv"
KEY_CSV = DATA / "plausibility_rating_key.csv"
XLSX = DATA / "plausibility_rating_sheet.xlsx"

ACTION_TEXT = {
    "increase_difficulty": "Make the next puzzle harder",
    "decrease_difficulty": "Make the next puzzle easier",
    "change_task_type": "Switch to a different type of puzzle",
    "offer_hint": "Show a hint straight away",
    "add_encouragement": "Give encouragement",
    "introduce_challenge": "Give a challenge puzzle",
    "repeat_simplified": "Give a simpler puzzle of the same type",
    "advance_topic": "Move on to a new topic (a bit harder)",
}
STATE_TEXT = {
    "rushing": "Finding it easy (under-challenged)",
    "on_task": "On task",
    "struggling": "Needs support",
}
KIND_TEXT = {
    "arithmetic": "Sum", "missing_number": "Missing number", "word_problem": "Word problem",
    "number_pattern": "Number pattern", "shape_pattern": "Shape pattern", "letter_pattern": "Letter pattern",
    "odd_one_out": "Odd one out",
}

INSTRUCTIONS = [
    ("How to rate", True),
    ("Each row is a real moment from the pilot at Excella School (21 October 2026), just after a learner answered a puzzle.", False),
    ("Look at what the learner did, then at what SensAI would have done next, and ask:", False),
    ("Would this have been an appropriate thing to do for this learner at this moment?", True),
    ("", False),
    ("1 = Inappropriate: a teacher would not do this", False),
    ("2 = Acceptable: reasonable, though not necessarily what I would choose", False),
    ("3 = What I would have done", False),
    ("", False),
    ("Some moments appear twice with different actions. Rate each row on its own.", False),
    ("You are not told which actions came from the model alone and which were changed by a safety rule; that is on purpose.", False),
    ("Rate quickly, from your first judgement. Use the Notes column only if something stands out.", False),
    ("Rate on your own: the teacher's ratings go in a separate copy, so the two can be compared.", False),
    ("", False),
    ("Columns", True),
    ("Seconds to answer: from the puzzle appearing to pressing Check.  Seconds before typing: until the first key press or choice.", False),
    ("Pauses: number of 20-second stretches with no typing.  What SensAI detected: the learner's state from their last five answers.", False),
]

HEADERS = ["Item", "Learner", "Puzzle no.", "Puzzle", "Type", "Level (1-5)", "Got it right?", "Seconds to answer",
           "Seconds before typing", "Hints used", "Pauses (20 s)", "What SensAI detected", "What SensAI would do next",
           "Your rating (1-3)", "Notes (optional)"]
WIDTHS = [6, 8, 9, 34, 15, 9, 10, 10, 11, 8, 9, 26, 34, 12, 30]


def build() -> Path:
    sheet = pd.read_csv(SHEET_CSV)
    bank = QuestionBank()
    wb = Workbook()
    ws = wb.active
    ws.title = "Rate here"
    thin = Side(style="thin", color="BBBBBB")
    header_fill = PatternFill("solid", fgColor="E6F4F1")
    rate_fill = PatternFill("solid", fgColor="FFF6D6")

    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = Border(bottom=thin)
    for _, r in sheet.iterrows():
        q = bank.get(PILOT_FORM_IDS[int(r.question) - 1])
        ws.append([
            int(r["item"]), r.learner, int(r.question), q.prompt, KIND_TEXT[r.kind], int(r.level),
            "Yes" if bool(r.correct) else "No", float(r.rt_s),
            None if pd.isna(r.first_input_s) else float(r.first_input_s), int(r.hints), int(r.pauses),
            STATE_TEXT[r.state], ACTION_TEXT[r.action_shown], None, None,
        ])
    for i, w in enumerate(WIDTHS, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    ws.row_dimensions[1].height = 32
    rate_col = HEADERS.index("Your rating (1-3)") + 1
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        row[rate_col - 1].fill = rate_fill
        row[rate_col - 1].alignment = Alignment(horizontal="center", vertical="top")
    letter = ws.cell(row=1, column=rate_col).column_letter
    dv = DataValidation(type="list", formula1='"1,2,3"', allow_blank=True, showDropDown=False,
                        errorTitle="Rating", error="Please choose 1, 2 or 3.")
    dv.add(f"{letter}2:{letter}{ws.max_row}")
    ws.add_data_validation(dv)
    ws.freeze_panes = "B2"

    info = wb.create_sheet("How to rate", 0)
    for text, bold in INSTRUCTIONS:
        info.append([text])
        info.cell(row=info.max_row, column=1).font = Font(bold=bold, size=12 if bold else 11)
    info.column_dimensions["A"].width = 120
    wb.active = 1
    wb.save(XLSX)
    return XLSX


def score(path: Path = XLSX) -> dict:
    ws = load_workbook(path, data_only=True)["Rate here"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    rating_i = HEADERS.index("Your rating (1-3)")
    ratings = pd.DataFrame({"item": [r[0] for r in rows], "rating": [r[rating_i] for r in rows]})
    ratings["rating"] = pd.to_numeric(ratings.rating, errors="coerce")
    key = pd.read_csv(KEY_CSV)
    sheet = pd.read_csv(SHEET_CSV)[["item", "learner", "question", "state"]]
    df = ratings.merge(key, on="item").merge(sheet, on="item")
    if df.rating.notna().sum() == 0:
        raise SystemExit("No ratings found yet.")

    # One moment = one learner at one question. Every moment has a "dqn" item
    # (the DQN's own action). Where a guardrail changed the action there is
    # also a "guarded" item; otherwise the DQN's action is the final one.
    moments = []
    for (_, _), g in df.groupby(["learner", "question"]):
        own = g[g.version == "dqn"].rating.iloc[0]
        final = g[g.version == "guarded"].rating.iloc[0] if (g.version == "guarded").any() else own
        moments.append({"state": g.state.iloc[0], "dqn_alone": own, "with_guardrails": final,
                        "guardrail_used": bool((g.version == "guarded").any())})
    m = pd.DataFrame(moments)
    dqn = m.dropna(subset=["dqn_alone"])
    with_guardrails = m.with_guardrails.dropna()
    rated = df.dropna(subset=["rating"])

    def summary(values) -> dict:
        v = pd.Series(values, dtype=float)
        return {"n": int(v.size), "plausible_rate (rated 2 or 3)": round(float((v >= 2).mean()), 3),
                "mean_rating": round(float(v.mean()), 2), "counts": {int(k): int(c) for k, c in v.value_counts().sort_index().items()}}

    result = {
        "rated_items": len(rated), "of": len(df),
        "moments": len(m), "moments_where_a_guardrail_changed_the_action": int(m.guardrail_used.sum()),
        "dqn_alone": summary(dqn.dqn_alone),
        "dqn_with_guardrails": summary(with_guardrails),
        "by_state (dqn alone)": {s: summary(g.dqn_alone) for s, g in dqn.groupby("state")},
        "by_state (with guardrails)": {s: summary(g.with_guardrails.dropna()) for s, g in m.groupby("state")},
        "note": "One rater (the researcher). Agreement with a second rater (weighted Cohen's kappa) needs the teacher's copy.",
    }
    out = ROOT / "analysis" / "results" / "plausibility.json"
    out.write_text(json.dumps(result, indent=2))
    return result


if __name__ == "__main__":
    command = sys.argv[1] if len(sys.argv) > 1 else "build"
    if command == "build":
        print("Wrote", build())
    else:
        print(json.dumps(score(), indent=2))
